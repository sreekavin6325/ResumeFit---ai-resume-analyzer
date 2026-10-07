"""Resume document parsing for PDF, DOCX, and TXT uploads."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import BinaryIO

from services.text_preprocessor import normalize_text


class ResumeParseError(ValueError):
    """Raised when a supported resume cannot be read."""


def _read_bytes(source: bytes | bytearray | str | Path | BinaryIO) -> bytes:
    if isinstance(source, (bytes, bytearray)):
        return bytes(source)
    if isinstance(source, (str, Path)):
        return Path(source).read_bytes()
    if hasattr(source, "getvalue"):
        return source.getvalue()
    if hasattr(source, "read"):
        position = source.tell() if hasattr(source, "tell") else None
        data = source.read()
        if position is not None and hasattr(source, "seek"):
            source.seek(position)
        return data
    raise ResumeParseError("Unsupported file object.")


def extract_text_from_pdf(source: bytes | bytearray | str | Path | BinaryIO) -> str:
    try:
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(_read_bytes(source)))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:  # parser errors vary by pypdf version
        raise ResumeParseError(f"Could not read the PDF: {exc}") from exc
    cleaned = normalize_text(text)
    if not cleaned:
        raise ResumeParseError("No selectable text was found. Try a text-based PDF or DOCX file.")
    return cleaned


def extract_text_from_docx(source: bytes | bytearray | str | Path | BinaryIO) -> str:
    try:
        from docx import Document

        document = Document(BytesIO(_read_bytes(source)))
        blocks = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                blocks.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(blocks)
    except Exception as exc:
        raise ResumeParseError(f"Could not read the DOCX file: {exc}") from exc
    cleaned = normalize_text(text)
    if not cleaned:
        raise ResumeParseError("No text was found in the DOCX file.")
    return cleaned


def extract_resume_text(source: bytes | bytearray | str | Path | BinaryIO, filename: str | None = None) -> str:
    name = filename or getattr(source, "name", "")
    suffix = Path(name).suffix.lower()
    if suffix == ".pdf":
        return extract_text_from_pdf(source)
    if suffix == ".docx":
        return extract_text_from_docx(source)
    if suffix == ".txt":
        try:
            return normalize_text(_read_bytes(source).decode("utf-8-sig"))
        except UnicodeDecodeError as exc:
            raise ResumeParseError("TXT resumes must use UTF-8 encoding.") from exc
    raise ResumeParseError("Supported resume formats are PDF, DOCX, and TXT.")

